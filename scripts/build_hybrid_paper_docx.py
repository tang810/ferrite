from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper_hybrid_multilayer_ferrite_SERF_MEG_revised.docx"
FIG_DIR = ROOT / "paper" / "figures_tie"


def set_run_font(run, size=10.5, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(doc, text="", style=None, align=None, size=10.5, bold=False, italic=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run_font(r, size=12 if level == 1 else 10.5, bold=True)
    return p


def add_caption(doc, text):
    p = add_para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_figure(doc, filename, caption, width=6.2):
    path = FIG_DIR / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=9, bold=True)
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            hdr[i].width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=8.5)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = Inches(widths[i])
    add_para(doc)
    return table


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)

    title = add_para(
        doc,
        "Modeling and Application-Oriented Optimization of Hybrid Multilayer Ferrite Magnetic Shields for SERF-MEG Magnetometers",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=16,
        bold=True,
    )
    title.paragraph_format.space_after = Pt(8)
    add_para(doc, "Author Name", align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
    add_para(doc, "", size=6)

    add_heading(doc, "Abstract", 1)
    add_para(
        doc,
        "SERF-based magnetoencephalography requires a low-frequency magnetic shielding environment with both high shielding factor and controlled magnetic-noise contribution. This paper upgrades the shield concept from a ferrite-only multilayer cylinder to a hybrid multilayer structure: flexible ferrite sheets provide a low-conductivity base, while an outer amorphous or nanocrystalline ribbon provides an additional high-permeability magnetic path. The four-layer ferrite base, with Rin = 100 mm, L = 200 mm, 0.15 mm sheet thickness, and 0.08 mm interlayer gap, is treated as a manufacturable validation candidate rather than a global optimum. FEM results give SFx = 3.63 and IntH2_total = 4.16e-6 for this ferrite-only base. Adding a 0.20 mm amorphous outer layer increases SFx to 8.55 and reduces IntH2_total to 6.28e-7, while a nanocrystalline outer layer under the present initial-permeability assumption increases SFx to 25.15 and reduces IntH2_total to 5.33e-8. Axial single-cap simulations show a separate tradeoff: a +Z nanocrystalline cap improves SFz from 2.41 to 2.91 but increases the IntH2 integral relative to the open ferrite baseline. Therefore, the simulated IntH2 trends are used as magnetic-noise-related indicators, not as proof of final magnetometer sensitivity. Prototype shielding-factor tests, compensation-coil nulling, measured material permeability/loss, and magnetometer PSD measurements are required before claiming sensitivity improvement.",
    )
    add_para(doc, "Keywords: SERF-MEG; magnetic shielding; ferrite sheet; nanocrystalline ribbon; amorphous ribbon; shielding factor; magnetic noise; finite-element method.", italic=True)

    add_heading(doc, "I. INTRODUCTION", 1)
    intro = [
        "Magnetic shielding is a central enabling component for optically pumped magnetometers and SERF-based biomagnetic measurements. The residual field at the sensor must be small enough for stable operation, and the shield itself must not introduce excessive low-frequency magnetic noise. Conventional high-permeability metallic shields, such as permalloy or mu-metal structures, can provide strong attenuation, but their cost, weight, forming process, and conductive magnetic-noise mechanisms motivate alternative structures for compact or flexible magnetometer platforms.",
        "Ferrite materials are attractive for low-noise shielding because of their high electrical resistivity and reduced eddy-current loss. However, the shielding factor of a thin flexible ferrite cylinder can be limited when the available radial space and total thickness are constrained. A ferrite-only design is therefore not sufficient as the final application structure if the target is both compactness and improved attenuation.",
        "This work investigates a hybrid multilayer shield in which flexible ferrite sheets form a low-conductivity base and an additional amorphous or nanocrystalline ribbon forms a high-permeability outer layer. Amorphous and nanocrystalline ribbons are commercially available soft-magnetic materials and are compatible with ribbon-based, conformal, and lower-cost shield fabrication. Compared with rigid bulk metallic shells, a ferrite-sheet plus ribbon strategy is more aligned with lightweight and potentially wearable magnetic shielding concepts.",
        "The paper deliberately separates simulated shielding improvement from experimental magnetometer sensitivity. The FEM model uses representative initial relative permeability values for low-field operation, because the shield is intended to work near the geomagnetic field without strong excitation. The actual material B-H slope near H = 0 and the magnetic loss term mu'' must be measured before the simulated noise-related integral can be converted into a quantitative magnetometer-noise prediction.",
        "The contributions are: (1) a manufacturable four-layer ferrite base is selected under total-thickness, radial-space, flexibility, assembly, SF, and IntH2 constraints; (2) amorphous and nanocrystalline outer layers are compared as hybrid enhancement candidates; (3) a single-ended axial cap is evaluated as an SFz and IntH2 tradeoff; and (4) an experimental validation route is defined using P0/P1/P2 prototypes, compensation-coil nulling, and magnetometer PSD measurements.",
    ]
    for t in intro:
        add_para(doc, t)
    add_figure(
        doc,
        "Fig01_concept_hybrid_multilayer_shield.png",
        "Fig. 1. Concept of the hybrid multilayer ferrite shield for magnetometer applications. (a) Four-layer ferrite base. (b) Outer amorphous or nanocrystalline high-permeability layer. (c) Single-ended cap and magnetometer placement.",
    )

    add_heading(doc, "II. PRINCIPLE OF HYBRID MULTILAYER MAGNETIC SHIELDING", 1)
    add_heading(doc, "A. Shielding Factor and Leakage Ratio", 2)
    add_para(doc, "The shielding factor is defined as SF = B0 / Bcenter, where B0 is the applied reference field and Bcenter is the field at the magnetometer position. For axial cases, the leakage ratio is reported as Bcenter / B0, so a lower leakage ratio corresponds to stronger attenuation.")
    add_heading(doc, "B. Magnetic-Noise-Related IntH2 Metric", 2)
    add_para(doc, "The FEM post-processing uses the field-energy-related integral IntH2 = integral |H|^2 dV over the magnetic materials as a geometry-dependent indicator connected to magnetic noise. For hybrid structures, the total integral must include both ferrite and added soft-magnetic ribbon regions. It should not be interpreted alone as an absolute noise prediction unless material loss is known.")
    add_para(doc, "A material-separated noise proxy is therefore written as")
    add_para(doc, "Noise proxy proportional to mu''_ferrite * IntH2_ferrite + mu''_metal * IntH2_metal.", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
    add_para(doc, "This expression is the reason that amorphous and nanocrystalline material characterization remains part of the validation route. A lower IntH2_total under equal-loss assumptions is favorable, but measured mu'' and PSD data are still required.")
    add_heading(doc, "C. Hybrid Outer-Layer Concept", 2)
    add_para(doc, "The ferrite base provides a low-conductivity and flexible shield body. The amorphous or nanocrystalline outer layer provides a higher-permeability flux path that improves transverse attenuation. In the current simulation stage, the amorphous and nanocrystalline layers are modeled with representative initial permeability assumptions of mu_r = 10000 and mu_r = 50000, respectively. These values represent low-field initial-permeability inputs and will be updated after B-H characterization.")
    add_heading(doc, "D. Axial Cap Concept", 2)
    add_para(doc, "Open-ended cylindrical shields generally leak in the axial direction. A single-ended cap is introduced to reduce axial leakage while preserving a practical opening. A small axial gap is included between the cylindrical shield and the cap to represent mechanical support or assembly spacing and to avoid an unrealistic zero-gap contact in the FEM mesh.")

    add_heading(doc, "III. FEM MODEL AND STRUCTURAL DESIGN MATRIX", 1)
    add_heading(doc, "A. Four-Layer Ferrite Base and Mechanical Constraints", 2)
    add_para(doc, "The selected ferrite base has Rin = 100 mm and L = 200 mm. Four ferrite sheets of thickness t = 0.15 mm are separated by g = 0.08 mm gaps, giving a total ferrite thickness of 0.60 mm and radial space occupation of 0.84 mm. The four-layer stack is not claimed as a global optimum. It is selected as a manufacturable validation candidate under the current total-thickness, radial-space, bending, assembly, SF, and IntH2 constraints.")
    add_figure(
        doc,
        "Fig02_ferrite_base_selection.png",
        "Fig. 2. Four-layer ferrite validation base. (a) Geometrical parameters of the selected manufacturable candidate. (b) Baseline shielding factor and IntH2 metric summary.",
    )
    add_heading(doc, "B. Design Matrix", 2)
    add_table(
        doc,
        ["Group", "Configuration", "Purpose", "Key modeling input"],
        [
            ["Ferrite base", "Four-layer ferrite cylinder", "Manufacturable low-noise base", "N = 4, t = 0.15 mm, g = 0.08 mm"],
            ["Hybrid AM", "Ferrite + amorphous outer layer", "Transverse SF enhancement reference", "outer t = 0.20 mm, gap = 0.40 mm, initial mu_r = 10000"],
            ["Hybrid NC", "Ferrite + nanocrystalline outer layer", "Priority simulated outer-layer candidate", "outer t = 0.20 mm, gap = 0.40 mm, initial mu_r = 50000"],
            ["P0", "Single amorphous shell + +Z amorphous cap", "Shielded reference, not unshielded baseline", "shell t = 0.20 mm, cap t = 0.20 mm"],
            ["Axial cap", "Four-layer ferrite + +Z nanocrystalline cap", "Axial leakage reduction tradeoff", "cap t = 0.20 mm, initial mu_r = 50000"],
        ],
        widths=[1.0, 2.0, 2.0, 2.7],
    )
    add_heading(doc, "C. Geomagnetic Residual-Field Scaling", 2)
    add_para(doc, "The FEM excitation field is approximately 1.2566 uT and is not itself a geomagnetic-field simulation. For engineering interpretation, residual fields under a 50 uT external geomagnetic field are estimated by linear scaling with the simulated SF. This is a linear initial-permeability estimate, not a nonlinear B-H geomagnetic-field simulation. Because the estimated residual fields remain at the uT level, compensation coils are required before magnetometer PSD testing.")

    add_heading(doc, "IV. SIMULATION RESULTS AND STRUCTURAL OPTIMIZATION", 1)
    add_heading(doc, "A. Four-Layer Ferrite Baseline", 2)
    add_para(doc, "For the transverse field direction, the four-layer ferrite-only base gives SFx = 3.629 and IntH2_total = 4.160e-6. This confirms that the ferrite stack is a practical low-conductivity base, but its attenuation is limited when used alone.")
    add_heading(doc, "B. Hybrid Outer-Layer Enhancement", 2)
    add_para(doc, "Compared with the ferrite-only base, the amorphous outer layer increases SFx from 3.63 to 8.55, while the nanocrystalline outer layer further increases SFx to 25.15. Meanwhile, IntH2_total decreases from 4.16e-6 to 6.28e-7 and 5.33e-8 for the amorphous and nanocrystalline hybrid cases, respectively. The integration includes both the ferrite layers and the added high-permeability layer.")
    add_table(
        doc,
        ["Case", "SFx", "Gain vs ferrite", "IntH2_total", "IntH2_outer", "Allowed interpretation"],
        [
            ["Ferrite only", "3.629", "1.00", "4.160e-6", "-", "Manufacturable baseline"],
            ["Ferrite + amorphous", "8.548", "2.36", "6.283e-7", "1.563e-7", "Improves SFx in simulation"],
            ["Ferrite + nanocrystalline", "25.146", "6.93", "5.333e-8", "1.117e-8", "Current preferred simulated outer layer"],
        ],
        widths=[1.8, 0.75, 1.0, 1.05, 1.05, 2.0],
    )
    add_figure(
        doc,
        "Fig03_hybrid_outer_layer_enhancement.png",
        "Fig. 3. Hybrid outer-layer enhancement in the transverse field direction. AM: amorphous; NC: nanocrystalline. (a) Shielding factor SFx. (b) Noise-related integral IntH2 total.",
    )
    add_heading(doc, "C. Axial Single-Cap Tradeoff", 2)
    add_para(doc, "The open four-layer ferrite cylinder gives SFz = 2.412 and leakage ratio = 0.4145. The P0 amorphous shell-cap reference gives SFz = 2.526 and IntH2_total = 3.960e-7. The four-layer ferrite plus single +Z nanocrystalline cap improves SFz to 2.907 and reduces the leakage ratio to 0.3440, but IntH2_total increases to 3.736e-6 relative to the open ferrite baseline. Thus the cap should be described as an axial SFz improvement with an IntH2 tradeoff, not as a proven magnetometer-sensitivity improvement.")
    add_table(
        doc,
        ["Axial case", "SFz", "Leakage ratio", "IntH2_total", "Paper claim"],
        [
            ["Open four-layer ferrite", "2.412", "0.4145", "2.971e-6", "Axial baseline"],
            ["P0 amorphous shell + +Z cap", "2.526", "0.3959", "3.960e-7", "Shielded reference for prototype route"],
            ["Ferrite + +Z nanocrystalline cap", "2.907", "0.3440", "3.736e-6", "Improves SFz but increases IntH2"],
        ],
        widths=[2.1, 0.75, 1.0, 1.1, 2.7],
    )
    add_figure(
        doc,
        "Fig04_axial_cap_tradeoff.png",
        "Fig. 4. Axial cap tradeoff. NC: nanocrystalline. (a) Axial shielding factor SFz. (b) IntH2 total. (c) Leakage ratio.",
    )
    add_heading(doc, "D. Geomagnetic Residual-Field Estimate", 2)
    add_para(doc, "Under linear scaling to a 50 uT external field, the estimated transverse residual field decreases from 13.78 uT for the ferrite-only base to 5.85 uT with the amorphous outer layer and 1.99 uT with the nanocrystalline outer layer. In the axial direction, the open ferrite case gives 20.73 uT, the P0 amorphous shell-cap reference gives 19.79 uT, and the four-layer ferrite plus nanocrystalline cap gives 17.20 uT. These residual fields are still too large for direct SERF operation, so three-axis compensation-coil nulling is part of the experimental procedure.")
    add_figure(
        doc,
        "Fig05_geomagnetic_residual_estimate.png",
        "Fig. 5. Estimated residual field under a 50 uT geomagnetic field using linear SF scaling. AM: amorphous; NC: nanocrystalline. (a) Transverse field. (b) Axial field.",
    )
    add_heading(doc, "E. Material-Loss Sensitivity", 2)
    add_para(doc, "The hybrid conclusion depends on both geometry and material loss. A sensitivity sweep of mu''_metal / mu''_ferrite shows how the non-ferrite layer can change the weighted noise proxy. Therefore, nanocrystalline is selected as the current simulated outer-layer priority because it gives the highest SFx and lowest IntH2_total under the assumed parameters, while final noise performance must be validated with measured mu'' and magnetometer PSD data.")
    add_figure(
        doc,
        "Fig06_material_loss_sensitivity.png",
        "Fig. 6. Sensitivity of the noise-amplitude proxy to the assumed magnetic loss of amorphous or nanocrystalline materials.",
    )

    add_heading(doc, "V. DISCUSSION AND EXPERIMENTAL VALIDATION ROUTE", 1)
    add_heading(doc, "A. Prototype SF Measurement", 2)
    add_para(doc, "The prototype plan uses three practical configurations: P0, a single amorphous shell plus single +Z amorphous cap; P1, a four-layer ferrite prototype; and P2, a four-layer ferrite prototype with a selected nanocrystalline outer layer. An optional two-layer ferrite prototype can be added to document the layer-number, flexibility, and assembly tradeoff. The required SF measurement should cover at least x and z directions.")
    add_heading(doc, "B. Compensation-Coil Assisted Residual-Field Nulling", 2)
    add_para(doc, "Because the linear 50 uT scaling predicts residual fields at the uT level, the magnetometer experiment requires three-axis compensation coils before PSD measurement. The relevant recorded quantities include residual DC field in nT and peak-to-peak drift in nT after nulling.")
    add_heading(doc, "C. Magnetometer PSD Measurement", 2)
    add_para(doc, "The final validation should report magnetometer PSD metrics, including noise_1Hz_fT_sqrtHz, noise_10Hz_fT_sqrtHz, and noise_avg_1_30Hz_fT_sqrtHz. Only after the P0/P1/P2 SF measurements, material permeability/loss measurements, compensation-coil nulling, and PSD results are available should the paper claim real magnetometer sensitivity improvement.")
    add_para(doc, "The long-term application target is a lightweight and potentially wearable magnetic shielding route for OPM- or SERF-based biomagnetic measurement. The present work supports this direction through modeling and prototype planning, but it does not claim that wearable MEG or final SERF sensitivity has already been achieved.")

    add_heading(doc, "VI. CONCLUSION", 1)
    add_para(doc, "A hybrid multilayer ferrite magnetic shield has been formulated for SERF-MEG magnetometer applications. The four-layer ferrite base is selected as a manufacturable validation candidate under thickness, flexibility, assembly, SF, and IntH2 constraints, not as a global optimum. In transverse simulations, adding a 0.20 mm amorphous outer layer increases SFx from 3.63 to 8.55, while a 0.20 mm nanocrystalline outer layer increases SFx to 25.15 and reduces IntH2_total to 5.33e-8 under the present initial-permeability assumptions. Therefore, the nanocrystalline outer layer is the current preferred simulated candidate. Axial single-cap simulations improve SFz but introduce an IntH2 tradeoff, so the cap design must be judged together with PSD experiments. The next step is fabrication and testing of P0/P1/P2 prototypes, including SF measurement, B-H based initial permeability and loss characterization, compensation-coil nulling, and magnetometer PSD validation.")

    for p in doc.paragraphs:
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(4)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
