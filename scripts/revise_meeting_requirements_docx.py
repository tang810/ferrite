from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "paper_intro_methods_20260528_V2_revised_clean.docx"
OUT = ROOT / "paper_intro_methods_20260528_V2_revised_clean_meeting_hybrid.docx"


def clear_paragraph(paragraph):
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def set_para(paragraph, text: str):
    style = paragraph.style
    alignment = paragraph.alignment
    clear_paragraph(paragraph)
    paragraph.style = style
    paragraph.alignment = alignment
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    return paragraph


def set_title(paragraph, text: str):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)
    run.bold = True


def insert_paragraph_after(paragraph, text: str = "", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    if style is not None:
        new_para.style = style
    if text:
        set_para(new_para, text)
    return new_para


def insert_table_after(doc: Document, paragraph, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    if doc.tables:
        table.style = doc.tables[0].style
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(text)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8)
                if r == 0:
                    for run in p.runs:
                        run.bold = True
    tbl = table._tbl
    paragraph._p.addnext(tbl)
    return table


def add_picture_after(doc: Document, paragraph, path: Path, width_in=6.2):
    p = insert_paragraph_after(paragraph, "")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    return p


def first_para(doc: Document, contains: str, start: int = 0):
    for i, p in enumerate(doc.paragraphs[start:], start):
        if contains in p.text:
            return i, p
    raise ValueError(f"Paragraph containing {contains!r} not found")


def append_case_rows(table):
    existing = {row.cells[0].text.strip() for row in table.rows}
    additions = [
        [
            "H_C2_N4_outer_AM",
            "4 + AM",
            "0.15 + 0.20",
            "0.08/0.40",
            "x",
            "hybrid amorphous outer-layer case",
        ],
        [
            "H_C2_N4_outer_NC",
            "4 + NC",
            "0.15 + 0.20",
            "0.08/0.40",
            "x",
            "hybrid nanocrystalline outer-layer case",
        ],
        [
            "P0_AM_shell_cap",
            "AM shell + cap",
            "0.20",
            "asm.",
            "z",
            "screened prototype reference, not no-shield",
        ],
        [
            "ZCAP_C2_N4_NC",
            "4 + NC cap",
            "0.15 + 0.20",
            "axial",
            "z",
            "axial single-cap tradeoff case",
        ],
    ]
    for row_values in additions:
        if row_values[0] in existing:
            continue
        row = table.add_row()
        for i, text in enumerate(row_values):
            row.cells[i].text = text
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8)


def cleanup_obsolete_paragraphs(doc: Document):
    obsolete_starts = [
        "IV. EXPERIMENTAL IMPROVEMENT OF SF AND MAGNETIC NOISE",
        "E.  Mesh and Boundary-Domain Convergence",
        "Mesh  convergence  data  are  available  for  C1",
        "F.  Axial Shielding",
        "The      no-shield      z-directed",
        "quired permeability comparison between",
        "I.  Field-Map Interpretation",
        "AEDT field-map exports",
        "This  ranking  is  directional",
        "x-directed  continuous-shell",
        "and the current geometry matrix.",
        "The IntH2 results should be read as magnetic-noise-related loss indicators, not measured magnetic noise. The layerwise",
        "needed to explain this distribution",
        "The material-parameter dependence is also unresolved.",
    ]
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if any(text.startswith(prefix) for prefix in obsolete_starts):
            set_para(paragraph, "")


def main():
    shutil.copy2(SRC, OUT)
    doc = Document(OUT)

    # Title, abstract, and keywords.
    set_title(
        doc.paragraphs[1],
        "Modeling and Application-Oriented Optimization of Hybrid Multilayer Ferrite Magnetic Shields for SERF-MEG Magnetometers",
    )
    set_para(
        doc.paragraphs[6],
        "Abstract--SERF-based magnetoencephalography requires a low-field magnetic environment that combines adequate shielding factor with controlled shield-induced magnetic-noise contribution. This paper revises the ferrite-only multilayer concept into a hybrid multilayer structure in which flexible ferrite sheets form a low-conductivity base and an additional amorphous or nanocrystalline ribbon provides a high-permeability magnetic path. The ferrite base has an inner radius of 100 mm and an axial length of 200 mm. A four-layer ferrite candidate with 0.15 mm sheets and 0.08 mm radial gaps is selected as a manufacturable validation base under total-thickness, radial-space, flexibility, assembly, shielding-factor, and IntH2 constraints, not as a global optimum.",
    )
    set_para(
        doc.paragraphs[7],
        "The FEM model uses directional no-shield references and component-based shielding factors. The validated transverse ferrite base gives SFx = 3.629 and IntH2_total = 4.160e-6. Adding an amorphous outer layer increases SFx to 8.548 and reduces IntH2_total to 6.283e-7, while a nanocrystalline outer layer increases SFx to 25.146 and reduces IntH2_total to 5.333e-8 under the present initial-permeability assumptions. Axial cap simulations show that a single nanocrystalline cap improves SFz from 2.412 to 2.907 but introduces an IntH2 tradeoff. Prototype SF measurement, low-field permeability/loss characterization, compensation-coil nulling, and magnetometer PSD measurement remain required before claiming instrument-level sensitivity improvement.",
    )
    set_para(
        doc.paragraphs[8],
        "Index Terms--Amorphous ribbon, ferrite magnetic shield, finite-element method, flexible magnetic shielding, magnetic noise, nanocrystalline ribbon, SERF-MEG, shielding factor.",
    )

    # Introduction.
    set_para(
        doc.paragraphs[11],
        "SERF-based MEG and other optically pumped magnetometer systems require a low-field and low-noise magnetic environment because femtotesla-level biomagnetic signals can be obscured by residual environmental fields and by magnetic noise generated inside the shield itself [1]-[5]. The shield is therefore not only a passive attenuation component. It must be evaluated by directional shielding factor, residual center field, and material-loss-related magnetic-noise indicators.",
    )
    set_para(
        doc.paragraphs[14],
        "Amorphous and nanocrystalline ribbons are commercially available high-permeability soft-magnetic materials. Compared with conventional permalloy or mu-metal shells, ribbon-based amorphous or nanocrystalline layers can be more compatible with low-cost and flexible shield fabrication. These engineering considerations are important for compact SERF-MEG or OPM-oriented shielding, where material availability, assembly, bending, and cost matter in addition to attenuation.",
    )
    set_para(
        doc.paragraphs[15],
        "Ferrite materials are attractive for low-noise shielding because of their high electrical resistivity and reduced eddy-current loss. However, the shielding factor of a thin flexible ferrite cylinder can be limited when the available radial space and total thickness are constrained. A purely metallic high-permeability shield can provide stronger attenuation, but may introduce cost, weight, eddy-current, and magnetic-loss concerns. A hybrid ferrite-metallic soft-magnetic structure is therefore investigated to improve attenuation while retaining the low-noise motivation of the ferrite base.",
    )
    set_para(
        doc.paragraphs[16],
        "In the proposed hybrid shield, flexible ferrite sheets form a low-conductivity multilayer base, and an additional amorphous or nanocrystalline ribbon forms a high-permeability outer magnetic path. The ferrite base is responsible for flexible low-conductivity shielding, while the outer ribbon is used to raise transverse shielding factor. The combination is intended as a manufacturable composite magnetic shielding route rather than a ferrite-only design.",
    )
    set_para(
        doc.paragraphs[17],
        "The present FEM stage uses representative constant relative permeability values as low-field initial-permeability assumptions. Since the shield is intended to operate near the geomagnetic field without a strong excitation field, the initial permeability can be interpreted as the slope of the B-H curve around H = 0. The amorphous and nanocrystalline permeability and loss parameters must be updated after material characterization, and the simulated IntH2 trends must be verified by prototype PSD measurements.",
    )
    set_para(
        doc.paragraphs[19],
        "The present study therefore evaluates a hybrid multilayer ferrite shield through a validation-gated evidence chain: a four-layer ferrite base is selected as a manufacturable candidate; amorphous and nanocrystalline outer layers are compared for transverse shielding enhancement; a single-ended cap is evaluated for axial leakage; residual fields are estimated under 50 uT linear scaling; and a P0/P1/P2 prototype route is defined for later SF, permeability, compensation-coil, and PSD validation.",
    )

    # Principle and methods.
    set_para(doc.paragraphs[30], "II. PRINCIPLE OF HYBRID MULTILAYER MAGNETIC SHIELDING")
    _, p34 = first_para(doc, "Two finite-element models are used")
    p = insert_paragraph_after(
        p34,
        "For the hybrid structure, the IntH2 integration includes both the ferrite sheets and the added amorphous or nanocrystalline outer layer. The total geometry-dependent indicator is separated as IntH2_total = IntH2_ferrite + IntH2_metal so that the material loss contribution can be discussed without assuming identical imaginary permeability for all materials.",
        p34.style,
    )
    p = insert_paragraph_after(
        p,
        "A material-separated noise proxy is therefore proportional to mu''_ferrite IntH2_ferrite + mu''_metal IntH2_metal. A lower IntH2_total is favorable under equal-loss assumptions, but it is not a direct proof of lower measured magnetic noise unless mu''(f), temperature, virtual-coil normalization, and magnetometer PSD data are also available.",
        p34.style,
    )
    p = insert_paragraph_after(
        p,
        "The amorphous and nanocrystalline ribbons are modeled by representative initial relative permeability values in the present low-field FEM stage. These values should be read as engineering assumptions for H approximately 0 operation and will be replaced by measured low-field B-H and loss data before final prototype claims.",
        p34.style,
    )
    set_para(doc.paragraphs[96], "III. FEM MODEL AND HYBRID STRUCTURAL DESIGN MATRIX")
    set_para(
        doc.paragraphs[118],
        "B. FEM Calculation Matrix and Data Export",
    )
    set_para(
        doc.paragraphs[119],
        "The calculation matrix is expanded from ferrite-only transverse continuous shells to a hybrid evidence chain. It includes the four-layer ferrite base, ferrite plus amorphous outer layer, ferrite plus nanocrystalline outer layer, a screened P0 amorphous shell-cap reference, and a four-layer ferrite case with a single nanocrystalline axial cap. The P0 case is intentionally not a no-shield reference, because the physical SERF residual field may be too large to measure without shielding.",
    )
    set_para(
        doc.paragraphs[120],
        "Each external-field case exports Bcenter components and the appropriate no-shield B0 reference. Each virtual pickup-coil case exports IntH2_total and material- or layer-resolved contributions, including IntH2_outer for hybrid outer layers. Missing values are marked as missing or pending and are not replaced by zero.",
    )
    if len(doc.tables) > 2:
        append_case_rows(doc.tables[2])
    set_para(
        doc.paragraphs[136],
        "E. Prototype Measurement and Mechanical-Test Route",
    )
    set_para(
        doc.paragraphs[137],
        "No prototype shielding-factor or magnetometer PSD result is claimed in the present manuscript. The planned validation route uses three practical configurations: P0, a single amorphous shell with a single +Z amorphous cap; P1, a four-layer ferrite prototype; and P2, a four-layer ferrite prototype with the selected nanocrystalline outer layer. Optional two-layer ferrite samples may be added to document bending and assembly limits.",
    )

    # Results.
    set_para(doc.paragraphs[139], "IV. SIMULATION RESULTS AND STRUCTURAL OPTIMIZATION")
    set_para(doc.paragraphs[140], "A. Four-Layer Ferrite Base and Hybrid Outer-Layer Enhancement")
    set_para(
        doc.paragraphs[141],
        "The x-directed no-shield reference gives B0,x = 1.256637e-6 T and passes the direction-purity check. The four-layer ferrite base C2_N4_t015_g008_x gives SFx = 3.629 and IntH2_total = 4.160e-6. This confirms the base as a manufacturable low-conductivity candidate, but the attenuation is limited if the ferrite stack is used alone. Compared with this ferrite-only base, the amorphous outer layer increases SFx to 8.548, while the nanocrystalline outer layer increases SFx to 25.146. Meanwhile, IntH2_total decreases to 6.283e-7 and 5.333e-8 for the amorphous and nanocrystalline hybrid cases, respectively. The IntH2 integration includes both the ferrite layers and the added outer high-permeability layer.",
    )
    set_para(doc.paragraphs[143], "B. Material-Separated Magnetic-Noise-Related Loss")
    set_para(
        doc.paragraphs[144],
        "The decrease in IntH2_total indicates a favorable geometry-dependent trend under the current material assumptions, but it does not by itself prove that measured shield noise will decrease. For hybrid structures, the relevant noise proxy depends on both ferrite and metallic-ribbon loss: mu''_ferrite IntH2_ferrite + mu''_metal IntH2_metal. The nanocrystalline outer layer is therefore identified as the current simulated priority candidate because it gives the highest SFx and the lowest IntH2_total, while final material-noise claims remain conditional on measured mu'' and PSD validation.",
    )
    set_para(doc.paragraphs[146], "C. Axial Single-Cap Tradeoff and Geomagnetic Residual Estimate")
    set_para(
        doc.paragraphs[147],
        "For axial shielding, the open four-layer ferrite cylinder gives SFz = 2.412 and LeakageRatioz approximately 0.4145. The P0 amorphous shell-cap reference gives SFz = 2.526 and IntH2_total = 3.960e-7. A four-layer ferrite shield with a single +Z nanocrystalline cap improves SFz to 2.907 and reduces the leakage ratio to 0.3440, but the total IntH2 increases relative to the open ferrite axial case. This is treated as an SFz-IntH2 tradeoff, not as a demonstrated magnetometer-sensitivity improvement. In the axial cap simulation, a small axial gap is introduced between the cylinder and cap to represent mechanical support and to avoid an unrealistically sharp zero-gap FEM contact.",
    )
    set_para(
        doc.paragraphs[172],
        "E. Mesh, Boundary-Domain, and Model-Use Caveats",
    )
    set_para(
        doc.paragraphs[173],
        "Mesh convergence data remain available for representative ferrite cases. These checks support scalar trend reporting but not a complete numerical-robustness claim because total element count, adaptive-pass count, and air-domain metadata are incomplete. Hybrid and cap results are therefore reported as current FEM evidence requiring continued numerical and prototype validation.",
    )
    set_para(
        doc.paragraphs[174],
        "F. Axial Shielding Characterization",
    )
    set_para(
        doc.paragraphs[175],
        "The z-directed no-shield reference has passed direction validation. Current axial results are reported for the open four-layer ferrite cylinder, the P0 amorphous shell-cap reference, and the four-layer ferrite case with a single nanocrystalline cap. These results are used only to describe an axial SFz-IntH2 tradeoff and are not generalized to a final closed-shield design.",
    )
    set_para(
        doc.paragraphs[170],
        "Layerwise IntH2 summation validation passed for all seven transverse continuous-shell cases. The multilayer fractions are listed in Table VI and plotted in Fig. 5. The four-layer case shows the largest fraction in the outermost layer under the present virtual-coil normalization. This is interpreted as a geometry-dependent integral distribution, not as a measured magnetic-noise spectrum.",
    )
    set_para(
        doc.paragraphs[179],
        "G. Segmented-Shell and Assembly Limits",
    )
    set_para(
        doc.paragraphs[180],
        "Segmented aligned and staggered models are still required to quantify manufacturable slot penalties. The continuous-shell results should therefore be read as idealized magnetic-shell evidence. The physical prototype plan must still check bending stiffness, assembly gaps, and whether four ferrite sheets are near the practical mounting limit.",
    )
    set_para(
        doc.paragraphs[181],
        "For the amorphous and nanocrystalline additions, the relative permeability values used here represent initial-permeability assumptions for low-field operation. They are placeholders for the slope of the measured B-H curve around H = 0 and must be updated after material characterization.",
    )
    set_para(
        doc.paragraphs[182],
        "The application-oriented interpretation is likewise conservative. The results suggest a potential route toward lightweight OPM/SERF shielding, but a practical system still requires P0/P1/P2 SF measurements, compensation-coil nulling, and magnetometer PSD validation.",
    )
    set_para(
        doc.paragraphs[202],
        "H. Initial-Permeability and Material-Loss Sensitivity",
    )
    set_para(
        doc.paragraphs[203],
        "The amorphous and nanocrystalline ribbons are modeled with representative initial relative permeability values for low-field operation. These constants are not arbitrary final material claims; they approximate the low-field B-H slope around H = 0 under geomagnetic-field operation. The model must be updated after measured permeability and magnetic-loss data are obtained.",
    )
    for idx in range(204, 209):
        set_para(doc.paragraphs[idx], "")
    set_para(
        doc.paragraphs[210],
        "I. Application-Oriented Interpretation",
    )
    set_para(
        doc.paragraphs[211],
        "Under linear scaling to a 50 uT external geomagnetic field, the estimated transverse residual field decreases from 13.78 uT for the ferrite-only base to 5.85 uT with the amorphous outer layer and 1.99 uT with the nanocrystalline outer layer. These values are linear initial-permeability estimates, not nonlinear geomagnetic B-H simulations.",
    )
    set_para(
        doc.paragraphs[212],
        "The long-term application target is a lightweight and potentially wearable magnetic shielding route for OPM- or SERF-based biomagnetic measurements. The present work supports this direction through modeling and prototype planning, but does not claim that wearable MEG or final SERF sensitivity has already been realized.",
    )
    set_para(doc.paragraphs[213], "")

    # Discussion and conclusion.
    set_para(
        doc.paragraphs[216],
        "The revised results show that the ferrite-only multilayer base is useful as a low-conductivity manufacturable platform, but hybridization is needed to raise the shielding factor toward an application-oriented range. The nanocrystalline outer layer gives the strongest simulated transverse attenuation and the lowest IntH2_total in the current dataset. This preference is conditional on the initial-permeability assumptions and must be revisited after measured permeability and loss data are available.",
    )
    set_para(
        doc.paragraphs[217],
        "The four-layer ferrite base is not claimed as a global optimum.",
    )
    set_para(
        doc.paragraphs[218],
        "Although additional ferrite sheets may further improve shielding, the multilayer stack becomes increasingly difficult to bend and assemble.",
    )
    set_para(
        doc.paragraphs[219],
        "The four-layer ferrite base is therefore selected as a manufacturable validation candidate under the present total-thickness, radial-space, flexibility, and assembly constraints. The P0/P1/P2 prototype plan is designed to test this engineering claim rather than to replace it with unsupported simulation-only conclusions.",
    )
    set_para(
        doc.paragraphs[220],
        "The IntH2 results should be read as magnetic-noise-related loss indicators, not measured magnetic noise. For hybrid structures, material loss must be separated into ferrite and metallic-ribbon terms. A lower IntH2_total under equal-loss assumptions is favorable, but the final sensitivity claim requires prototype SF measurements, measured mu''(f), compensation-coil nulling, and magnetometer PSD data.",
    )
    set_para(
        doc.paragraphs[231],
        "The application claim is also conservative. Linear scaling to a 50 uT geomagnetic field indicates residual fields at the uT level, so three-axis compensation coils are required before SERF operation. Therefore, the paper presents a potential route toward compact OPM/SERF shielding rather than a completed wearable MEG demonstration.",
    )
    for idx in range(232, 236):
        set_para(doc.paragraphs[idx], "")
    set_para(
        doc.paragraphs[238],
        "A hybrid multilayer ferrite magnetic shield has been formulated for SERF-MEG magnetometer applications. The four-layer ferrite base is selected as a manufacturable validation candidate under thickness, flexibility, assembly, SF, and IntH2 constraints, not as a global optimum. In transverse simulations, the ferrite-only base gives SFx = 3.629 and IntH2_total = 4.160e-6. Adding an amorphous outer layer increases SFx to 8.548, and adding a nanocrystalline outer layer increases SFx to 25.146 while reducing IntH2_total to 5.333e-8 under the present initial-permeability assumptions.",
    )
    set_para(
        doc.paragraphs[239],
        "Axial cap simulations show that a single nanocrystalline cap improves SFz but introduces an IntH2 tradeoff. The current conclusions remain conditional: the amorphous and nanocrystalline permeability values represent low-field initial-permeability assumptions, material loss must be measured, and prototype P0/P1/P2 SF and PSD tests are required before claiming instrument-level sensitivity improvement.",
    )
    set_para(doc.paragraphs[240], "")
    set_para(
        doc.paragraphs[244],
        "The prototype validation plan uses P0, a single amorphous shell plus single +Z amorphous cap; P1, a four-layer ferrite prototype; and P2, a four-layer ferrite prototype with a selected nanocrystalline outer layer. The required measurements include x- and z-directed SF, low-field permeability and loss characterization, compensation-coil residual-field nulling, and magnetometer PSD.",
    )

    cleanup_obsolete_paragraphs(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
